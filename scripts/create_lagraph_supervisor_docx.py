from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import html


PROJECT_ROOT = Path(r"D:\la_v12")
HTML_PATH = PROJECT_ROOT / "docs" / "lagraph_supervisor_report_2026_06_12.html"
DOCX_PATH = PROJECT_ROOT / "docs" / "lagraph_supervisor_report_2026_06_12.docx"


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=12.5, bold=False, color="111827"):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def text_content(node):
    return " ".join(" ".join(node.itertext()).split())


def add_text_with_inline(paragraph, node, base_size=12.5):
    if node.text:
        run = paragraph.add_run(node.text)
        set_run_font(run, size=base_size)
    for child in node:
        tag = child.tag.lower() if isinstance(child.tag, str) else ""
        child_text = "".join(child.itertext())
        if tag in {"b", "strong"}:
            run = paragraph.add_run(child_text)
            set_run_font(run, size=base_size, bold=True)
        elif tag in {"code"}:
            run = paragraph.add_run(child_text)
            set_run_font(run, name="Consolas", size=base_size - 0.5, color="334155")
        elif tag in {"br"}:
            paragraph.add_run("\n")
        else:
            add_text_with_inline(paragraph, child, base_size=base_size)
        if child.tail:
            run = paragraph.add_run(child.tail)
            set_run_font(run, size=base_size)


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {min(level, 3)}"
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    if level == 1:
        set_run_font(run, size=20, bold=True, color="1F4D78")
    elif level == 2:
        set_run_font(run, size=16, bold=True, color="2E5E8D")
    else:
        set_run_font(run, size=13.5, bold=True, color="334155")
    return paragraph


def add_paragraph_from_node(doc, node, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.18
    add_text_with_inline(paragraph, node)
    return paragraph


def add_note_paragraph(doc, node):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.05)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    add_text_with_inline(paragraph, node)
    set_paragraph_font(paragraph, size=12.2)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2E74B5")
    p_bdr.append(left)
    p_pr.append(p_bdr)
    return paragraph


def add_list_item(doc, text, ordered=False):
    paragraph = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=12.3)


def add_table_from_node(doc, table_node):
    rows = table_node.xpath(".//tr")
    if not rows:
        return
    max_cols = max(len(row.xpath("./th|./td")) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    for r_idx, row_node in enumerate(rows):
        cells = row_node.xpath("./th|./td")
        for c_idx in range(max_cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if c_idx < len(cells):
                text = text_content(cells[c_idx])
            else:
                text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            set_run_font(run, size=9.2 if max_cols >= 6 else 10.5, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
    doc.add_paragraph()


def add_code_block(doc, node):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    text = node.text_content().strip()
    run = paragraph.add_run(text)
    set_run_font(run, name="Consolas", size=10.5, color="111827")


def add_metadata_table(doc, root):
    meta_items = []
    for item in root.xpath(".//header//div[contains(@class,'meta')]/div"):
        label_nodes = item.xpath("./b")
        label = text_content(label_nodes[0]) if label_nodes else ""
        value = text_content(item).replace(label, "", 1).strip()
        meta_items.append((label, value))
    if not meta_items:
        return
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    for label, value in meta_items:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
        set_cell_shading(row[0], "F2F4F7")
        for idx, cell in enumerate(row):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=10.5, bold=(idx == 0))
    doc.add_paragraph()


def process_node(doc, node):
    tag = node.tag.lower() if isinstance(node.tag, str) else ""
    cls = node.get("class", "")
    if tag == "h2":
        add_heading(doc, text_content(node), 1)
    elif tag == "h3":
        add_heading(doc, text_content(node), 2)
    elif tag == "p":
        add_paragraph_from_node(doc, node)
    elif tag == "pre":
        add_code_block(doc, node)
    elif tag == "ul":
        for li in node.xpath("./li"):
            add_list_item(doc, text_content(li), ordered=False)
    elif tag == "ol":
        for li in node.xpath("./li"):
            add_list_item(doc, text_content(li), ordered=True)
    elif tag == "table":
        add_table_from_node(doc, node)
    elif tag == "div" and any(name in cls.split() for name in ["note", "risk", "ok", "warn"]):
        add_note_paragraph(doc, node)
    elif tag == "div" and "card" in cls.split():
        heading = text_content(node.xpath("./h3")[0]) if node.xpath("./h3") else ""
        if heading:
            add_heading(doc, heading, 3)
        for child in node:
            if child.tag.lower() not in {"h3", "span"}:
                process_node(doc, child)
    elif tag == "div":
        for child in node:
            process_node(doc, child)
    elif tag == "section":
        for child in node:
            process_node(doc, child)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(12.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(12.3)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("LaGraph 项目汇报与论文问题梳理")
    set_run_font(run, size=9.5, color="666666")


def main():
    html_text = HTML_PATH.read_text(encoding="utf-8")
    root = html.fromstring(html_text)

    doc = Document()
    configure_document(doc)

    title = text_content(root.xpath("//header/h1")[0])
    subtitle = text_content(root.xpath("//header/p[contains(@class,'subtitle')]")[0])
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(3)
    run = title_p.add_run(title)
    set_run_font(run, size=24, bold=True, color="111827")

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    sub_run = sub_p.add_run(subtitle)
    set_run_font(sub_run, size=13.5, color="555555")

    add_metadata_table(doc, root)

    add_heading(doc, "目录", 1)
    for idx, link in enumerate(root.xpath("//nav//ol/li/a"), start=1):
        add_list_item(doc, f"{idx}. {text_content(link)}", ordered=False)
    doc.add_page_break()

    for section in root.xpath("//main/section"):
        for child in section:
            process_node(doc, child)

    add_footer(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
